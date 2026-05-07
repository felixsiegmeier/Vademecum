import asyncio
import csv
import io
import json
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Literal, Optional

from fastapi import Body, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from dotenv import load_dotenv
from openai import APIConnectionError, APIStatusError, RateLimitError
from pydantic import BaseModel, field_validator
from workflows.brief.verlauf import validate_curate_variant as _validate_curate_variant
from storage import brief_storage, learning_storage
from skills import learning
from tools.patient_tools import TOOL_ARGS, TOOL_FUNCTIONS
from workflows.document_extraction.orchestrator import extract_proposals, extract_proposals_streaming
from utils.tool_loop import Proposal
from workflows.patient_chat.orchestrator import CHAT_2PASS_CUTOFF, run_single_pass_chat
from workflows.stammdaten_extraction.orchestrator import extract_stammdaten
from workflows.meilenstein import orchestrator as meilenstein_orchestrator
from llm_client import LLMClient
from models.patient import Patient, PatientSummary, Stammdaten
from storage.patients import (
    delete_patient,
    list_patient_ids,
    load_meilenstein,
    load_patient,
    next_patient_id,
    patient_yaml_hash,
    save_meilenstein,
    delete_meilenstein,
    save_patient,
    update_meilenstein_content,
)

from workflows.brief import orchestrator as brief
from storage.brief_storage import BRIEF_SECTIONS as _BRIEF_SECTIONS
from utils.prompts import get_prompt as _get_prompt

load_dotenv()

# Binäre MIME-Typen (PDF + Bilder) — auch für /api/extract-stammdaten akzeptiert
_BINARY_UPLOAD_MIMES: frozenset[str] = frozenset({
    "application/pdf",
    "image/jpeg",
    "image/jpg",
    "image/png",
    "image/gif",
    "image/webp",
    "image/tiff",
    "image/bmp",
})

_TEXT_MIMES: frozenset[str] = frozenset({
    "text/plain",
    "text/markdown",
})
_CSV_MIMES: frozenset[str] = frozenset({
    "text/csv",
    "application/csv",
})
_XLSX_MIMES: frozenset[str] = frozenset({
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})
_DOCX_MIMES: frozenset[str] = frozenset({
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
})

# /api/uploads akzeptiert alle Formate; /api/extract-stammdaten nur binäre
ALLOWED_UPLOAD_MIMES: frozenset[str] = (
    _BINARY_UPLOAD_MIMES | _TEXT_MIMES | _CSV_MIMES | _XLSX_MIMES | _DOCX_MIMES
)


def _bytes_to_text(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1")


def _csv_to_markdown(file_bytes: bytes) -> str:
    text = _bytes_to_text(file_bytes)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return text
    header = rows[0]
    sep = ["---"] * len(header)
    lines = (
        ["| " + " | ".join(header) + " |", "| " + " | ".join(sep) + " |"]
        + ["| " + " | ".join(row) + " |" for row in rows[1:]]
    )
    return "\n".join(lines)


def _xlsx_to_markdown(file_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [[str(c) if c is not None else "" for c in row] for row in ws.values]
        if not rows:
            continue
        parts.append(f"## {sheet_name}")
        sep = ["---"] * len(rows[0])
        parts.append("| " + " | ".join(rows[0]) + " |")
        parts.append("| " + " | ".join(sep) + " |")
        parts.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        parts.append("")
    return "\n".join(parts)


def _docx_to_text(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []
    for block in doc.element.body:
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            text = Paragraph(block, doc).text.strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            from docx.table import Table
            tbl = Table(block, doc)
            tbl_rows = [[c.text.replace("\n", " ") for c in row.cells] for row in tbl.rows]
            if tbl_rows:
                parts.append("| " + " | ".join(tbl_rows[0]) + " |")
                parts.append("| " + " | ".join(["---"] * len(tbl_rows[0])) + " |")
                parts.extend("| " + " | ".join(r) + " |" for r in tbl_rows[1:])
            parts.append("")
    return "\n".join(parts)

app = FastAPI(title="arztbrief-app")
# Ein globaler LLM-Client für alle Requests (hält die HTTP-Session offen)
llm = LLMClient()

# Erlaubt Requests vom Vite-Dev-Server (localhost:5173) und späteren Deployments
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ── Lazy-Loader für System-Prompts ────────────────────────────────────────────
# ── Request/Response-Modelle ──────────────────────────────────────────────────

class Message(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    messages: list[Message]


class PatientChatRequest(BaseModel):
    messages: list[dict[str, Any]]


class CreateStammdatenRequest(BaseModel):
    name: str
    geburtsdatum: date
    geschlecht: Literal["m", "w", "d"]
    bettplatz: str
    aufnahmedatum: date
    aufnahme_quelle: Optional[Literal["elektiv", "notfall", "extern"]] = None
    verlegungsziel: Optional[str] = None


class CreatePatientRequest(BaseModel):
    stammdaten: CreateStammdatenRequest


class ApplyToolsRequest(BaseModel):
    calls: list[dict]


class ApplyProposalsRequest(BaseModel):
    proposals: list[Proposal]
    force: bool = False


class AktivRequest(BaseModel):
    aktiv: bool


class MeilensteinGenerateRequest(BaseModel):
    current_meilenstein: Optional[str] = None


class MeilensteinUpdateRequest(BaseModel):
    content: str


class LearnFromEditsRequest(BaseModel):
    patient_id: str
    edited_content: str


class SaveRulesRuleItem(BaseModel):
    section: str
    rule_text: str


class SaveRulesRequest(BaseModel):
    rules_to_add: list[SaveRulesRuleItem] = []
    rule_ids_to_delete: list[str] = []


class RebuildRuleRequest(BaseModel):
    section: str
    original_rule_text: str
    original_reasoning: str = ""
    anchor: str = ""
    clarification: str


class BriefAgentGenerateRequest(BaseModel):
    extra_context: str = ""
    adressat: Optional[str] = None
    curate_variant: Optional[str] = None

    @field_validator("curate_variant")
    @classmethod
    def _check_curate_variant(cls, v: Optional[str]) -> Optional[str]:
        if v is None:
            return v
        return _validate_curate_variant(v)


# ── Patienten-CRUD ────────────────────────────────────────────────────────────

@app.get("/api/health")
async def health():
    return {"status": "ok"}


@app.get("/api/patients")
def get_patients():
    """Gibt eine schlanke Liste aller Patienten zurück (nur für die Sidebar)."""
    patients = [load_patient(pid) for pid in list_patient_ids()]
    summaries = [
        PatientSummary(
            id=p.stammdaten.id,
            name=p.stammdaten.name,
            aktiv=p.stammdaten.aktiv,
            bettplatz=p.stammdaten.bettplatz,
            aufnahmedatum=p.stammdaten.aufnahmedatum,
        )
        for p in patients
    ]
    return JSONResponse(
        content=[s.model_dump(mode="json") for s in summaries],
        media_type="application/json; charset=utf-8",
    )


@app.get("/api/patients/{patient_id}")
def get_patient(patient_id: str):
    """Gibt den vollständigen Patient-Datensatz zurück (für PatientLanding)."""
    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail=f"Patient {patient_id} not found")
    return JSONResponse(
        content=patient.model_dump(mode="json"),
        media_type="application/json; charset=utf-8",
    )


@app.post("/api/patients")
def create_patient(req: CreatePatientRequest):
    """Legt einen neuen Patienten an und gibt die generierte ID zurück."""
    patient_id = next_patient_id()
    s = req.stammdaten
    patient = Patient(
        stammdaten=Stammdaten(
            id=patient_id,
            name=s.name,
            geburtsdatum=s.geburtsdatum,
            geschlecht=s.geschlecht,
            bettplatz=s.bettplatz,
            aufnahmedatum=s.aufnahmedatum,
            aufnahme_quelle=s.aufnahme_quelle,
            verlegungsziel=s.verlegungsziel,
            aktiv=True,
        ),
    )
    save_patient(patient)
    return JSONResponse(
        content={"patient_id": patient_id},
        media_type="application/json; charset=utf-8",
    )


@app.delete("/api/patients/{patient_id}")
def delete_patient_endpoint(patient_id: str):
    """Löscht Patient + alle zugehörigen Dateien (YAML, Meilenstein, Brief)."""
    try:
        load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")
    delete_patient(patient_id)
    return JSONResponse(content={"ok": True})


@app.patch("/api/patients/{patient_id}/aktiv")
def set_aktiv(patient_id: str, req: AktivRequest):
    """Setzt den aktiv/inaktiv-Status (steuert Sidebar-Sichtbarkeit)."""
    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")
    patient.stammdaten.aktiv = req.aktiv
    save_patient(patient)
    return JSONResponse(content={"ok": True, "aktiv": req.aktiv})


# ── Meilenstein ───────────────────────────────────────────────────────────────
# Der Meilenstein ist eine vom LLM generierte Markdown-Übersicht des Patientenstatus.
# Er wird separat gespeichert (.md + .meta.json) und kann vom Arzt bearbeitet werden.
# is_stale = True, wenn das YAML nach der letzten Generierung geändert wurde.

@app.get("/api/patients/{patient_id}/meilenstein")
def get_meilenstein(patient_id: str):
    try:
        load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    result = load_meilenstein(patient_id)
    if result is None:
        raise HTTPException(404, "Kein Meilenstein vorhanden")

    md_content, meta = result
    # Hash des aktuellen YAML vergleichen mit dem Hash zum Zeitpunkt der Generierung
    current_hash = patient_yaml_hash(patient_id)
    is_stale = meta.get("yaml_hash", "") != current_hash

    return JSONResponse(content={
        "content": md_content,
        "generated_at": meta.get("generated_at"),
        "is_stale": is_stale,
    })


@app.post("/api/patients/{patient_id}/meilenstein/generate")
async def generate_meilenstein(
    patient_id: str,
    req: Optional[MeilensteinGenerateRequest] = Body(default=None),
):
    """LLM generiert Meilenstein aus dem aktuellen Patient-YAML.

    Wenn req.current_meilenstein gefüllt: Konsolidierungs-Modus (manuelle Änderungen bleiben).
    Wenn leer/null: Generations-Modus (reine YAML-Extraktion).
    """
    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    current_meilenstein = (req.current_meilenstein if req else None) or ""

    try:
        md_content = await meilenstein_orchestrator.generate(
            patient,
            llm,
            current_meilenstein=current_meilenstein,
        )
    except RateLimitError:
        raise HTTPException(429, "LLM Rate-Limit erreicht. Bitte kurz warten.")
    except APIConnectionError as e:
        raise HTTPException(503, f"LLM nicht erreichbar: {e}")
    except APIStatusError as e:
        raise HTTPException(502, f"LLM API Fehler {e.status_code}: {e.message}")

    yaml_hash = patient_yaml_hash(patient_id)
    generated_at = datetime.now(timezone.utc).isoformat()
    meta = {"yaml_hash": yaml_hash, "generated_at": generated_at}
    save_meilenstein(patient_id, md_content, meta)

    return JSONResponse(content={
        "content": md_content,
        "generated_at": generated_at,
        "is_stale": False,
    })

@app.put("/api/patients/{patient_id}/meilenstein")
def update_meilenstein(patient_id: str, req: MeilensteinUpdateRequest):
    """Speichert manuelle Bearbeitungen des Meilensteins (ohne Neuberechnung des Hashes)."""
    try:
        load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")
    update_meilenstein_content(patient_id, req.content)
    return JSONResponse(content={"ok": True})

@app.delete("/api/patients/{patient_id}/meilenstein")
def delete_meilenstein_endpoint(patient_id: str):
    meilenstein = load_meilenstein(patient_id)
    if meilenstein is None:
        raise HTTPException(404, "Kein Meilenstein vorhanden")
    delete_meilenstein(patient_id)
    return JSONResponse(content={"ok": True})


# ── Learn — Meilenstein ───────────────────────────────────────────────────────

@app.post("/api/learn/meilenstein/from-edits")
async def learn_meilenstein_from_edits(req: LearnFromEditsRequest):
    """Vergleicht editierten Meilenstein mit dem generierten Original und extrahiert Lernregeln."""
    last_generated = learning_storage.load_last_output(req.patient_id, domain="meilenstein")
    if last_generated is None:
        raise HTTPException(404, "Kein generierter Meilenstein als Referenz vorhanden.")

    if req.edited_content.strip() == last_generated.strip():
        return JSONResponse(content={"rule_candidates": [], "trivial_changes": []})

    existing_rules = learning_storage.load_rules(domain="meilenstein")
    result_candidates, trivial_changes = await learning.from_edits(
        llm, last_generated, req.edited_content, existing_rules
    )
    return JSONResponse(content={"rule_candidates": result_candidates, "trivial_changes": trivial_changes})


@app.post("/api/learn/meilenstein/save-rules")
def learn_meilenstein_save_rules(req: SaveRulesRequest):
    existing = learning_storage.load_rules(domain="meilenstein")

    deleted_count = 0
    if req.rule_ids_to_delete:
        ids_to_delete = set(req.rule_ids_to_delete)
        before = len(existing)
        existing = [r for r in existing if r.id not in ids_to_delete]
        deleted_count = before - len(existing)

    saved_count = 0
    for item in req.rules_to_add:
        try:
            rule = learning_storage.new_rule(section=item.section, rule_text=item.rule_text)
        except Exception as exc:
            raise HTTPException(422, f"Ungültige Regel: {exc}")
        existing.append(rule)
        saved_count += 1

    learning_storage.save_rules(existing, domain="meilenstein")
    return JSONResponse(content={
        "saved_count": saved_count,
        "deleted_count": deleted_count,
        "total_rules": len(existing),
    })


@app.post("/api/learn/meilenstein/rebuild-rule")
async def learn_meilenstein_rebuild_rule(req: RebuildRuleRequest):
    result = await learning.rebuild(
        client=llm,
        section=req.section,
        original_rule_text=req.original_rule_text,
        original_reasoning=req.original_reasoning,
        anchor=req.anchor,
        clarification=req.clarification,
    )
    return JSONResponse(content={
        "section": req.section,
        "rule_text": result.rule_text,
        "reasoning": result.reasoning,
        "anchor": req.anchor,
    })


@app.get("/api/learn/meilenstein/system-prompt")
def learn_meilenstein_system_prompt():
    prompt = meilenstein_orchestrator._load_prompt()
    return JSONResponse(content={"content": prompt})


@app.get("/api/learn/meilenstein/rules")
def learn_meilenstein_rules():
    rules = learning_storage.load_rules(domain="meilenstein")
    return JSONResponse(content={"rules": [r.model_dump() for r in rules]})


@app.delete("/api/learn/meilenstein/rules/{rule_id}", status_code=204)
def learn_meilenstein_delete_rule(rule_id: str):
    existing = learning_storage.load_rules(domain="meilenstein")
    filtered = [r for r in existing if r.id != rule_id]
    if len(filtered) == len(existing):
        raise HTTPException(404, f"Regel {rule_id} nicht gefunden.")
    learning_storage.save_rules(filtered, domain="meilenstein")


# ── Learn — Brief ─────────────────────────────────────────────────────────────

_BRIEF_SECTION_PROMPT_FILES: dict[str, str] = {
    "diagnosen": "prompt.md",
    "anamnese": "prompt.md",
    "therapie": "prompt.md",
}

_BRIEF_SECTION_PROMPT_DIRS: dict[str, Path] = {
    "diagnosen": Path(__file__).parent / "workflows" / "brief" / "diagnosen",
    "anamnese": Path(__file__).parent / "workflows" / "brief" / "anamnese",
    "therapie": Path(__file__).parent / "workflows" / "brief" / "therapie",
    "verlauf": Path(__file__).parent / "workflows" / "brief" / "verlauf",
}


def _assert_valid_brief_section(section: str) -> None:
    if section not in learning_storage.BRIEF_SECTIONS_WITH_LEARNING:
        raise HTTPException(404, f"Sektion '{section}' nicht lernfähig.")


def _get_brief_section_system_prompt(section: str) -> str:
    if section == "verlauf":
        # Für Display/Learning-Kontext: shared + kompakt als repräsentativer Default.
        curate_dir = _BRIEF_SECTION_PROMPT_DIRS["verlauf"] / "03_curate" / "prompts"
        shared = _get_prompt("shared.md", curate_dir)
        specific = _get_prompt("kompakt.md", curate_dir)
        return shared + "\n\n" + specific
    filename = _BRIEF_SECTION_PROMPT_FILES[section]
    prompt_dir = _BRIEF_SECTION_PROMPT_DIRS[section]
    return _get_prompt(filename, prompt_dir)


@app.post("/api/learn/brief/{section}/from-edits")
async def learn_brief_from_edits(section: str, req: LearnFromEditsRequest):
    _assert_valid_brief_section(section)
    last_generated = learning_storage.load_last_output(req.patient_id, domain="brief", section=section)
    if last_generated is None:
        raise HTTPException(404, "Kein generierter Output als Referenz vorhanden.")

    if req.edited_content.strip() == last_generated.strip():
        return JSONResponse(content={"rule_candidates": [], "trivial_changes": []})

    existing_rules = learning_storage.load_rules(domain="brief", section=section)
    result_candidates, trivial_changes = await learning.from_edits(
        llm, last_generated, req.edited_content, existing_rules
    )
    return JSONResponse(content={"rule_candidates": result_candidates, "trivial_changes": trivial_changes})


@app.post("/api/learn/brief/{section}/save-rules")
def learn_brief_save_rules(section: str, req: SaveRulesRequest):
    _assert_valid_brief_section(section)
    existing = learning_storage.load_rules(domain="brief", section=section)

    deleted_count = 0
    if req.rule_ids_to_delete:
        ids_to_delete = set(req.rule_ids_to_delete)
        before = len(existing)
        existing = [r for r in existing if r.id not in ids_to_delete]
        deleted_count = before - len(existing)

    saved_count = 0
    for item in req.rules_to_add:
        try:
            rule = learning_storage.new_rule(section=item.section, rule_text=item.rule_text)
        except Exception as exc:
            raise HTTPException(422, f"Ungültige Regel: {exc}")
        existing.append(rule)
        saved_count += 1

    learning_storage.save_rules(existing, domain="brief", section=section)
    return JSONResponse(content={
        "saved_count": saved_count,
        "deleted_count": deleted_count,
        "total_rules": len(existing),
    })


@app.post("/api/learn/brief/{section}/rebuild-rule")
async def learn_brief_rebuild_rule(section: str, req: RebuildRuleRequest):
    _assert_valid_brief_section(section)
    result = await learning.rebuild(
        client=llm,
        section=req.section,
        original_rule_text=req.original_rule_text,
        original_reasoning=req.original_reasoning,
        anchor=req.anchor,
        clarification=req.clarification,
    )
    return JSONResponse(content={
        "section": req.section,
        "rule_text": result.rule_text,
        "reasoning": result.reasoning,
        "anchor": req.anchor,
    })


@app.get("/api/learn/brief/{section}/system-prompt")
def learn_brief_system_prompt(section: str):
    _assert_valid_brief_section(section)
    prompt = _get_brief_section_system_prompt(section)
    return JSONResponse(content={"content": prompt})


@app.get("/api/learn/brief/{section}/rules")
def learn_brief_rules(section: str):
    _assert_valid_brief_section(section)
    rules = learning_storage.load_rules(domain="brief", section=section)
    return JSONResponse(content={"rules": [r.model_dump() for r in rules]})


@app.delete("/api/learn/brief/{section}/rules/{rule_id}", status_code=204)
def learn_brief_delete_rule(section: str, rule_id: str):
    _assert_valid_brief_section(section)
    existing = learning_storage.load_rules(domain="brief", section=section)
    filtered = [r for r in existing if r.id != rule_id]
    if len(filtered) == len(existing):
        raise HTTPException(404, f"Regel {rule_id} nicht gefunden.")
    learning_storage.save_rules(filtered, domain="brief", section=section)


@app.get("/api/learn/brief/{section}/last/{patient_id}")
def learn_brief_last_snapshot(section: str, patient_id: str):
    """Liefert den zuletzt gespeicherten Snapshot für eine Brief-Sektion. text=null wenn keiner."""
    _assert_valid_brief_section(section)
    text = learning_storage.load_last_output(patient_id, domain="brief", section=section)
    return JSONResponse(content={"text": text})


@app.get("/api/learn/meilenstein/last/{patient_id}")
def learn_meilenstein_last_snapshot(patient_id: str):
    """Liefert den zuletzt gespeicherten Meilenstein-Snapshot. text=null wenn keiner."""
    text = learning_storage.load_last_output(patient_id, domain="meilenstein")
    return JSONResponse(content={"text": text})


# ── Brief-Agent V1 ────────────────────────────────────────────────────────────
# Pfad: /api/brief/{patient_id}/*

@app.get("/api/brief/{patient_id}")
async def get_brief_agent(patient_id: str):
    """Liefert aktuellen Brief-State (alle 5 Sektionen). Leeres Skelett wenn noch kein Brief."""
    return brief_storage.load_brief(patient_id)


@app.delete("/api/brief/{patient_id}", status_code=204)
def delete_brief_agent_endpoint(patient_id: str):
    """Löscht Brief-Datei + last-Snapshots für alle Brief-Sektionen. Lernlog-Regeln bleiben."""
    if not brief_storage.delete_brief(patient_id):
        raise HTTPException(404, "Kein Brief vorhanden.")
    for section in learning_storage.BRIEF_SECTIONS_WITH_LEARNING:
        snap = learning_storage.SNAPSHOTS_DIR / "brief" / f"{patient_id}.yml"
        if snap.exists():
            snap.unlink()
            break  # eine Datei pro Patient, nicht pro Sektion


@app.post("/api/brief/{patient_id}/generate")
async def generate_brief_agent(
    patient_id: str,
    req: Optional[BriefAgentGenerateRequest] = Body(default=None),
):
    """Generiert Diagnosen + Anamnese + Therapie parallel, dann Verlauf sequenziell.
    Befunde-Sektion bleibt unverändert. extra_context ist ephemer (nicht persistiert)."""
    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    extra_context = req.extra_context if req else ""
    adressat = (req.adressat if req else None) or "normalstation_intern"
    curate_variant = req.curate_variant if req else None
    meilenstein_result = load_meilenstein(patient_id)
    meilenstein_text = meilenstein_result[0] if meilenstein_result else None

    current = brief_storage.load_brief(patient_id)
    befunde_existing = current.get("befunde", "")

    diag, anam, ther = await asyncio.gather(
        brief.generate_diagnosen(patient, extra_context=extra_context),
        brief.generate_anamnese(patient, extra_context=extra_context),
        brief.generate_therapie(patient, extra_context=extra_context),
    )
    verlauf = await brief.generate_verlauf(
        patient, meilenstein_text, befunde_existing, diag, anam, ther,
        extra_context=extra_context,
        adressat=adressat,
        curate_variant_override=curate_variant,
    )

    new_brief = {**current, "diagnosen": diag, "anamnese": anam, "therapie": ther, "verlauf": verlauf}
    brief_storage.save_brief(patient_id, new_brief)
    return brief_storage.load_brief(patient_id)


@app.post("/api/brief/{patient_id}/generate-section/{section}")
async def regenerate_section_agent(
    patient_id: str,
    section: str,
    req: Optional[BriefAgentGenerateRequest] = Body(default=None),
):
    """Re-generiert eine einzelne Sektion (nicht befunde). extra_context ephemer."""
    if section not in {"diagnosen", "anamnese", "therapie", "verlauf"}:
        raise HTTPException(400, f"Section '{section}' nicht re-generierbar.")

    adressat = req.adressat if req else None
    curate_variant = req.curate_variant if req else None
    if section != "verlauf" and (adressat is not None or curate_variant is not None):
        raise HTTPException(422, "adressat/curate_variant nur für section=verlauf gültig.")

    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    extra_context = req.extra_context if req else ""
    current = brief_storage.load_brief(patient_id)

    if section == "diagnosen":
        result = await brief.generate_diagnosen(patient, extra_context=extra_context)
    elif section == "anamnese":
        result = await brief.generate_anamnese(patient, extra_context=extra_context)
    elif section == "therapie":
        result = await brief.generate_therapie(patient, extra_context=extra_context)
    else:  # verlauf
        meilenstein_result = load_meilenstein(patient_id)
        meilenstein_text = meilenstein_result[0] if meilenstein_result else None
        result = await brief.generate_verlauf(
            patient,
            meilenstein_text,
            current.get("befunde", ""),
            current.get("diagnosen", ""),
            current.get("anamnese", ""),
            current.get("therapie", ""),
            extra_context=extra_context,
            adressat=adressat or "normalstation_intern",
            curate_variant_override=curate_variant,
        )

    brief_storage.update_section(patient_id, section, result)
    return {section: result}


@app.post("/api/brief/{patient_id}/format-befunde")
async def format_befunde_agent(patient_id: str, body: dict):
    """Pre-Pass: formatiert SAP-Roh-Befunde und persistiert in befunde-Sektion.
    body: {raw_text: str}. Returns: {befunde: <formatiert>}."""
    raw = body.get("raw_text", "").strip()
    if not raw:
        raise HTTPException(400, "raw_text leer.")
    extra_context = body.get("extra_context", "")
    formatted = await brief.format_sap_befunde(raw, extra_context=extra_context)
    brief_storage.update_section(patient_id, "befunde", formatted)
    return {"befunde": formatted}


@app.put("/api/brief/{patient_id}/section/{section}")
async def save_section_edit_agent(patient_id: str, section: str, body: dict):
    """Autosave für User-Edits ohne LLM. body: {content: str}. Returns: {<section>: <content>}."""
    if section not in _BRIEF_SECTIONS:
        raise HTTPException(400, f"Unknown section '{section}'.")
    content = body.get("content", "")
    brief_storage.update_section(patient_id, section, content)
    return {section: content}


@app.post("/api/brief/{patient_id}/polish-section/{section}")
async def polish_section_agent(
    patient_id: str,
    section: str,
    req: Optional[BriefAgentGenerateRequest] = Body(default=None),
):
    """Lektor-Korrektur einer einzelnen Sektion (alle 4). Kein Inhalt wird verändert.
    body: {extra_context?: str}. Returns: {<section>: <polished>}."""
    if section not in {"diagnosen", "anamnese", "therapie", "verlauf"}:
        raise HTTPException(400, f"Section '{section}' nicht polierbar.")

    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    extra_context = req.extra_context if req else ""
    current = brief_storage.load_brief(patient_id)
    current_text = current.get(section, "")
    if not current_text:
        raise HTTPException(400, f"Sektion '{section}' ist leer — bitte zuerst generieren.")

    result = await brief.polish_section(
        section=section,
        current_text=current_text,
        extra_context=extra_context,
        patient=patient,
    )
    brief_storage.update_section(patient_id, section, result)
    return {section: result}


@app.post("/api/extract-text")
async def extract_text(files: list[UploadFile] = File(...)):
    """Extrahiert Text aus einem oder mehreren Dateien.
    Text-Dateien werden direkt dekodiert; Binärdateien (PDF, Bilder) via LLM.
    Response: {combined_text: str}."""
    from llm_client import file_to_content_parts

    parts: list[str] = []
    for f in files:
        mime = f.content_type or "application/octet-stream"
        data = await f.read()
        if mime in _TEXT_MIMES:
            parts.append(_bytes_to_text(data))
        elif mime in _CSV_MIMES:
            parts.append(_csv_to_markdown(data))
        elif mime in _XLSX_MIMES:
            parts.append(_xlsx_to_markdown(data))
        elif mime in _DOCX_MIMES:
            parts.append(_docx_to_text(data))
        else:
            content_parts = file_to_content_parts(data, mime)
            resp = await brief._lite().chat_completion(
                [{"role": "user", "content": content_parts + [
                    {"type": "text", "text": "Gib den gesamten Text dieses Dokuments wieder. Keine Zusammenfassung, kein Kommentar."}
                ]}],
                temperature=0,
                max_tokens=4096,
            )
            parts.append((resp.choices[0].message.content or "").strip())

    return {"combined_text": "\n\n".join(p for p in parts if p)}


# ── Chat-History (persistente Text-Messages pro Patient) ─────────────────────

from storage.chat_storage import delete_chat, load_chat, save_chat
from models.chat import ChatHistory


@app.get("/api/chat/{patient_id}")
async def get_patient_chat_history(patient_id: str):
    """Gibt die gespeicherten Text-Chat-Nachrichten eines Patienten zurück."""
    return load_chat(patient_id).model_dump()


@app.put("/api/chat/{patient_id}")
async def put_patient_chat_history(patient_id: str, body: ChatHistory):
    """Überschreibt die Chat-History eines Patienten vollständig."""
    save_chat(patient_id, body)
    return {"ok": True}


@app.delete("/api/chat/{patient_id}", status_code=204)
async def delete_patient_chat_history(patient_id: str):
    """Löscht die Chat-History eines Patienten."""
    delete_chat(patient_id)


# ── Allgemeiner Chat (kein Patient-Kontext) ───────────────────────────────────

@app.post("/api/chat")
async def chat(request: ChatRequest):
    """Einfacher Chat ohne Patient-Kontext — für den allgemeinen Assistenten-Modus."""
    messages = [m.model_dump() for m in request.messages]

    # Server-Sent Events: jedes Text-Token wird sofort ans Frontend gestreamt
    async def event_stream():
        try:
            async for chunk in llm.stream_chat(messages):
                yield f"data: {chunk}\n\n"
            yield "data: [DONE]\n\n"
        except RateLimitError:
            yield "data: [ERROR] Rate-Limit erreicht. Bitte kurz warten.\n\n"
        except APIStatusError as e:
            yield f"data: [ERROR] LLM API Fehler {e.status_code}\n\n"
        except Exception as e:
            yield f"data: [ERROR] {e}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


# ── Patienten-Chat mit 2-Pass-Extraktion ─────────────────────────────────────

@app.post("/api/patients/{patient_id}/chat")
async def patient_chat(patient_id: str, req: PatientChatRequest):
    """
    Arzt chattet über einen Patienten.

    Kurze Inputs (≤ CHAT_2PASS_CUTOFF Zeichen): Single-Pass — LLM entscheidet
    selbst ob Tool-Call (→ proposals) oder Text-Antwort (→ reply).

    Lange Inputs (> CHAT_2PASS_CUTOFF): 2-Pass-Pipeline wie Upload-Endpoint.
    Typisch für eingefügte Akte-Texte. Gibt nur proposals + auto_skipped zurück.
    """
    try:
        patient = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    user_text = ""
    for m in reversed(req.messages):
        if m.get("role") == "user":
            content = m.get("content", "")
            if isinstance(content, str):
                user_text = content
            break
    if not user_text:
        raise HTTPException(400, "Keine Benutzernachricht gefunden")

    try:
        if len(user_text) > CHAT_2PASS_CUTOFF:
            # Langer Input → 2-Pass-Pipeline (wie Upload)
            proposals = await extract_proposals(llm, patient, user_text, content_type="text")
            if not proposals:
                return JSONResponse(
                    content={
                        "proposals": [],
                        "auto_skipped": True,
                        "message": "Keine Änderungen vorgeschlagen.",
                        "reply": None,
                    },
                    media_type="application/json; charset=utf-8",
                )
            return JSONResponse(
                content={
                    "proposals": [p.model_dump() for p in proposals],
                    "auto_skipped": False,
                    "reply": None,
                },
                media_type="application/json; charset=utf-8",
            )
        else:
            # Kurzer Input → Single-Pass (Tool-Call oder Text-Antwort)
            proposals, reply = await run_single_pass_chat(llm, patient, user_text, date.today())
            if reply:
                return JSONResponse(
                    content={
                        "proposals": [],
                        "auto_skipped": False,
                        "reply": reply,
                    },
                    media_type="application/json; charset=utf-8",
                )
            if proposals:
                return JSONResponse(
                    content={
                        "proposals": [p.model_dump() for p in proposals],
                        "auto_skipped": False,
                        "reply": None,
                    },
                    media_type="application/json; charset=utf-8",
                )
            return JSONResponse(
                content={
                    "proposals": [],
                    "auto_skipped": True,
                    "message": "Keine Änderungen vorgeschlagen.",
                    "reply": None,
                },
                media_type="application/json; charset=utf-8",
            )
    except RateLimitError:
        raise HTTPException(429, "LLM Rate-Limit erreicht. Bitte kurz warten und erneut versuchen.")
    except APIConnectionError as e:
        raise HTTPException(503, f"LLM nicht erreichbar: {e}")
    except APIStatusError as e:
        raise HTTPException(502, f"LLM API Fehler {e.status_code}: {e.message}")


# ── Dokument-Upload / Extraktion ──────────────────────────────────────────────

_PDF_MIMES = frozenset({"application/pdf"})
_IMAGE_MIMES = _BINARY_UPLOAD_MIMES - _PDF_MIMES


@app.post("/api/uploads")
async def upload_document(
    file: UploadFile = File(...),
    patient_id: str | None = Form(None),
    extra_context: str = Form(""),
):
    """
    Streaming-Upload-Endpoint (NDJSON, chunked HTTP).
    Analysiert ein Dokument per 2-Pass-Extraktion und streamt Events zurück:
    status → heartbeats → proposals (pro Iteration) → done (oder error).
    patient_id optional: Patientenstand als Kongruenz-Kontext.

    NDJSON statt SSE: POST-Upload, EventSource (nur GET/HEAD) nicht verwendbar.
    """
    mime_type = file.content_type or "application/octet-stream"
    if mime_type not in ALLOWED_UPLOAD_MIMES:
        raise HTTPException(415, f"Nicht unterstützter Dateityp: {mime_type}")

    file_bytes = await file.read()

    patient = None
    if patient_id:
        try:
            patient = load_patient(patient_id)
        except FileNotFoundError:
            raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    if mime_type in _PDF_MIMES:
        content_type: Literal["pdf", "image", "text"] = "pdf"
        file_content: str | bytes = file_bytes
    elif mime_type in _IMAGE_MIMES:
        content_type = "image"
        file_content = file_bytes
    elif mime_type in _TEXT_MIMES:
        content_type = "text"
        file_content = _bytes_to_text(file_bytes)
    elif mime_type in _CSV_MIMES:
        content_type = "text"
        file_content = _csv_to_markdown(file_bytes)
    elif mime_type in _XLSX_MIMES:
        content_type = "text"
        file_content = _xlsx_to_markdown(file_bytes)
    elif mime_type in _DOCX_MIMES:
        content_type = "text"
        file_content = _docx_to_text(file_bytes)
    else:
        raise HTTPException(415, f"Nicht unterstützter Dateityp: {mime_type}")

    async def ndjson_stream():
        async for event in extract_proposals_streaming(
            llm, patient, file_content, content_type, image_mime_type=mime_type,
            extra_context=extra_context,
        ):
            yield json.dumps(event, ensure_ascii=False) + "\n"

    return StreamingResponse(
        ndjson_stream(),
        media_type="application/x-ndjson",
        # Prevent Nginx from buffering the stream in reverse-proxy setups (e.g. Charité)
        headers={"X-Accel-Buffering": "no"},
    )



@app.post("/api/extract-stammdaten")
async def extract_stammdaten_endpoint(file: UploadFile = File(...)):
    """
    Single-Pass LLM-Extraktion der Patientenstammdaten aus einem Dokument.
    Alle Felder sind nullable — nicht erkannte Felder kommen als null.
    Kein Error bei Nicht-Patientendokumenten, nur leeres Ergebnis.
    """
    mime_type = file.content_type or "application/octet-stream"
    if mime_type not in _BINARY_UPLOAD_MIMES:
        raise HTTPException(400, f"Nicht unterstützter Dateityp: {mime_type}")

    file_bytes = await file.read()

    try:
        result = await extract_stammdaten(llm, file_bytes, mime_type)
    except RateLimitError:
        raise HTTPException(429, "LLM Rate-Limit erreicht. Bitte kurz warten.")
    except APIConnectionError as e:
        raise HTTPException(503, f"LLM nicht erreichbar: {e}")
    except APIStatusError as e:
        raise HTTPException(502, f"LLM API Fehler {e.status_code}: {e.message}")

    return JSONResponse(
        content=result.model_dump(mode="json"),
        media_type="application/json; charset=utf-8",
    )


@app.post("/api/patients/{patient_id}/apply-tools")
async def apply_tools(patient_id: str, req: ApplyToolsRequest):
    """
    Führt eine Liste von Tool-Calls direkt aus — ohne LLM-Beteiligung.
    Wird aufgerufen, wenn der Arzt Upload-Proposals bestätigt.
    Gibt pro Call das Ergebnis zurück (ok + summary oder error).
    """
    try:
        load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    results = []
    for call in req.calls:
        tool = call.get("tool")
        args = call.get("args", {})

        if tool not in TOOL_FUNCTIONS:
            results.append({"tool": tool, "ok": False, "error": f"Unknown tool: {tool}"})
            continue

        fn = TOOL_FUNCTIONS[tool]
        try:
            result = fn(patient_id, TOOL_ARGS[tool].model_validate(args))
            results.append({"tool": tool, **result})
        except TypeError as e:
            results.append({"tool": tool, "ok": False, "error": f"Argument error: {e}"})
        except Exception as e:
            results.append({"tool": tool, "ok": False, "error": str(e)})

    return JSONResponse(
        content={"results": results},
        media_type="application/json; charset=utf-8",
    )


_IDENTITY_FIELDS: frozenset[str] = frozenset({"name", "geburtsdatum", "geschlecht"})


def _detect_mismatch(patient: Patient, proposals: list[Proposal]) -> list[dict]:
    """Checks update_stammdaten proposals against current patient identity fields.

    Returns a list of {feld, current, proposed} for fields where the current
    value is non-empty and differs from the proposed value. Empty list = no conflict.
    When multiple proposals target the same field, the last one wins.
    """
    proposed: dict[str, object] = {}
    for p in proposals:
        if (
            p.type == "update_singleton"
            and p.call is not None
            and p.call.tool == "update_stammdaten"
        ):
            feld = p.call.args.get("feld")
            if feld in _IDENTITY_FIELDS:
                proposed[str(feld)] = p.call.args.get("wert")

    if not proposed:
        return []

    conflicts = []
    for feld, proposed_raw in proposed.items():
        current_raw = getattr(patient.stammdaten, feld, None)
        if current_raw is None:
            continue
        current_str = str(current_raw)
        if not current_str:
            continue
        proposed_str = str(proposed_raw) if proposed_raw is not None else ""
        if current_str != proposed_str:
            conflicts.append({"feld": feld, "current": current_str, "proposed": proposed_str})

    return conflicts


def _run_tool(patient_id: str, tool: str, args: dict) -> dict:
    """Führt ein einzelnes Tool aus und gibt das Ergebnis zurück."""
    fn = TOOL_FUNCTIONS.get(tool)
    if fn is None:
        return {"ok": False, "error": f"Unknown tool: {tool}"}
    try:
        return fn(patient_id, TOOL_ARGS[tool].model_validate(args))
    except TypeError as e:
        return {"ok": False, "error": f"Argument error: {e}"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def _to_apply_result(proposal_type: str, run_result: dict, summary_override: str | None = None) -> dict:
    """Normalisiert das Tool-Ergebnis auf das ApplyResult-Schema (immer `summary`, nie `error`)."""
    out: dict[str, Any] = {"type": proposal_type, "ok": run_result.get("ok", False)}
    if "id" in run_result:
        out["id"] = run_result["id"]
    out["summary"] = summary_override or run_result.get("summary") or run_result.get("error", "")
    return out


# delete_entry meldet eine fehlende ULID mit dieser Fehlermeldung — bei "update"
# behandeln wir das als Erfolg (gewünschter Endzustand: alter weg, neuer da).
_DELETE_ID_NOT_FOUND_MARKER = "Keine Liste enthält Eintrag mit id="


@app.post("/api/patients/{patient_id}/apply-proposals")
async def apply_proposals_endpoint(patient_id: str, req: ApplyProposalsRequest):
    """
    Führt eine Liste von Proposals aus — versteht Update-Gruppen (delete + add).

    Bei type="update" wird ZUERST `add_call` ausgeführt, danach `delete_call`.
    Schlägt add fehl → kein delete (Datenverlust ausgeschlossen).
    Schlägt delete fehl, weil die ID schon weg ist → wird als Erfolg gewertet.
    Schlägt delete aus anderen Gründen fehl → ok=false mit Hinweis auf manuelle Prüfung.
    """
    try:
        patient_snapshot = load_patient(patient_id)
    except FileNotFoundError:
        raise HTTPException(404, f"Patient {patient_id} nicht gefunden")

    if not req.force:
        conflicts = _detect_mismatch(patient_snapshot, req.proposals)
        if conflicts:
            return JSONResponse(
                status_code=409,
                content={"mismatch_warning": True, "conflicting_fields": conflicts},
                media_type="application/json; charset=utf-8",
            )

    results = []
    for proposal in req.proposals:
        if proposal.type == "update":
            if proposal.delete_call is None or proposal.add_call is None:
                results.append({
                    "type": "update",
                    "ok": False,
                    "summary": "Fehlende delete_call oder add_call",
                })
                continue
            add_result = _run_tool(patient_id, proposal.add_call.tool, proposal.add_call.args)
            if not add_result.get("ok"):
                results.append(_to_apply_result("update", add_result))
                continue
            del_result = _run_tool(patient_id, proposal.delete_call.tool, proposal.delete_call.args)
            if del_result.get("ok"):
                results.append(_to_apply_result("update", add_result))
                continue
            err = del_result.get("error", "") or ""
            if _DELETE_ID_NOT_FOUND_MARKER in err:
                results.append(_to_apply_result("update", add_result))
            else:
                id_suffix = (proposal.delete_call.args.get("id") or "")[-6:]
                results.append(_to_apply_result(
                    "update",
                    {"ok": False, "id": add_result.get("id")},
                    summary_override=(
                        f"Update teilweise — Eintrag {id_suffix} konnte nicht entfernt werden, "
                        f"manuell prüfen."
                    ),
                ))
        else:
            if proposal.call is None:
                results.append({"type": proposal.type, "ok": False, "summary": "Fehlende call"})
                continue
            run_result = _run_tool(patient_id, proposal.call.tool, proposal.call.args)
            results.append(_to_apply_result(proposal.type, run_result))

    return JSONResponse(
        content={"results": results},
        media_type="application/json; charset=utf-8",
    )
